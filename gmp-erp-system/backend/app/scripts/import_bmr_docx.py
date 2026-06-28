"""Импорт BMR/ЗПС из .docx фабричного формата NOVUGEN в структурированный
шаблон (BmrTemplate + BmrSection). Документ регулярный: повторяющиеся блоки
— `Процесс: … / Комната …` → тех.этапы; «ОТЧЁТ … ОКРУЖАЮЩЕЙ СРЕДЫ» → среда;
«ОБОРУДОВАНИЕ/ПРИБОРЫ» → оборудование; «ПРОИЗВОДСТВЕННАЯ ФОРМУЛА» → формула;
«ЛИСТ РАСПОЛОЖЕНИЯ» → лист распределения. Блоки часто повторяются ×3 по
рынкам — последовательные дубликаты схлопываем.

Результат — ЧЕРНОВИК шаблона (status=draft) для проверки технологом/ДОК
перед утверждением (СОП-11). Запуск:
    python -m app.scripts.import_bmr_docx /tmp/bmr.docx SIT-MET "НовуСита-М …"
"""
from __future__ import annotations

import re
import sys
from difflib import SequenceMatcher
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from app.core.database import SessionLocal
from app.models.identity import User
from app.models.inventory import BmrSection, BmrTemplate, Product
from app.models.master_data import Material


def _clean(s: str | None) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("\x00", " ")).strip()


def _short_label(s: str, limit: int = 80) -> str:
    """Метка поля — короткая. Длинные ячейки-инструкции/расчёты не годятся как
    подпись поля (полный текст всё равно остаётся в ячейке таблицы и виден при
    заполнении). Режем по границе слова и добавляем многоточие."""
    s = _clean(re.sub(r"_+", " ", s))
    # убрать повторяющиеся сегменты («… · (г или кг) · (г или кг)»)
    segs: list[str] = []
    for seg in (p.strip() for p in s.split("·")):
        if seg and (not segs or segs[-1] != seg):
            segs.append(seg)
    s = " · ".join(segs)
    if len(s) <= limit:
        return s
    cut = s[:limit].rsplit(" ", 1)[0] or s[:limit]
    return f"{cut.rstrip(' ·,;:')}…"


def _section_title(raw: str | None, fallback: str) -> str:
    title = _clean(raw)
    if not title or not re.search(r"[A-Za-zА-Яа-я0-9]", title):
        return fallback
    return title


def _sig_fields() -> list[dict]:
    return [
        {"label": "Выполнено · ДП", "type": "signature_operator"},
        {"label": "Проверено · ДОК", "type": "signature_qa"},
    ]


def _process_fields() -> list[dict]:
    return [
        {"label": "Дата-время начала", "type": "datetime"},
        {"label": "Дата-время окончания", "type": "datetime"},
        {"label": "Предыдущий ЛС", "type": "text"},
        {"label": "Предыдущая серия", "type": "text"},
    ]


def _norm_key(raw: str | None) -> str:
    text = (raw or "").lower().replace("ё", "е")
    text = re.sub(r"\*+", " ", text)
    text = text.replace("(", " ").replace(")", " ")
    text = re.sub(r"[^a-zа-я0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _material_alias_key(name: str | None) -> str:
    key = _norm_key(name)
    aliases = [
        (("метформин", "гидрохлорид"), "API-MET"),
        (("ситаглиптин", "фосфат"), "API-SITA"),
        (("повидон",), "EXC-POV"),
        (("лаурил", "сульфат"), "EXC-SLS"),
        (("кроскармеллоз",), "EXC-SOD"),
        (("коллоид", "кремн"), "EXC-AEROSIL"),
        (("aerosil",), "EXC-AEROSIL"),
        (("микрокристаллическ", "целлюлоз"), "EXC-MCC"),
        (("pharmasel",), "EXC-MCC"),
        (("стеарат", "магни"), "EXC-MGST"),
        (("opadry", "blue"), "EXC-OPA-BLUE"),
        (("opadry",), "COAT-OPADRY"),
        (("очищенная", "вода"), "UTIL-WATER"),
    ]
    for needles, code in aliases:
        if all(n in key for n in needles):
            return code
    return ""


def _field_pack(name: str) -> list[dict]:
    return [
        {"label": f"{name} · № серии сырья", "type": "text"},
        {"label": f"{name} · № аналит. листа", "type": "text"},
        {"label": f"{name} · вес нетто", "type": "number", "unit": "кг"},
        {"label": f"{name} · Выдал (Склад)", "type": "signature_warehouse"},
        {"label": f"{name} · Проверил (ДП)", "type": "signature_operator"},
        {"label": f"{name} · Проверил (ДОК)", "type": "signature_qa"},
    ]


def _distribution_fields(groups: list[dict]) -> list[dict]:
    fields: list[dict] = []
    for group in groups:
        for item in group.get("items", []):
            fields.extend(_field_pack(item.get("name") or "Материал"))
    return fields


def _iter_blocks(doc: Document):
    """Идёт по телу в порядке: ('p', text) | ('t', table)."""
    paras = {p._p: p for p in doc.paragraphs}
    tbls = {t._tbl: t for t in doc.tables}
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = paras.get(child)
            if p is not None and _clean(p.text):
                yield ("p", _clean(p.text))
        elif child.tag == qn("w:tbl"):
            t = tbls.get(child)
            if t is not None:
                yield ("t", t)


def _rows(t) -> list[list[str]]:
    return [[_clean(c.text) for c in r.cells] for r in t.rows]


def _cell_text(cell) -> str:
    return _clean(" ".join(p.text for p in cell.paragraphs))


def _strip_signature_text(text: str) -> str:
    # Бумажные строки подписи в конце шага: «Выполнил: ДП (Подпись и Дата)
    # Проверил: ДОК (Подпись и Дата)», «Подтвержден ДОК ____» — в eBMR их
    # заменяют кнопки э-подписи.
    text = re.sub(r"\s*(Выполнил|Выполнено|Проверил|Проверено|Подтвержден[оа]?)\s*:?\s*\(?\s*(ДП|ДОК)\b.*$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*\(\s*Подпись\s*(и|/)\s*Дата\s*\)\s*:?", "", text, flags=re.IGNORECASE)
    return _clean(text)


_BLANK_RE = re.compile(r"_{3,}")


def _text_blank_unit(after: str) -> tuple[str | None, str]:
    """Единица измерения по тексту сразу после прочерка → (unit, type)."""
    m = re.match(r"\s*(%|кг|мин\b|г\b|см\b|мл\b|°c)", after, re.IGNORECASE)
    if not m:
        return None, "text"
    return m.group(1).rstrip("."), "number"


def _extract_text_blanks(text: str, fields: list[dict], context: str) -> list[dict] | None:
    """Прочерки в тексте инструкции («Содержание влаги: ___% по массе») —
    это места ввода данных: на каждый создаём поле; фронт рендерит ввод
    прямо в строке текста (text_blanks по порядку прочерков)."""
    if not _BLANK_RE.search(text):
        return None
    parts = _BLANK_RE.split(text)
    blanks: list[dict] = []
    for i in range(len(parts) - 1):
        before = parts[i]
        after = parts[i + 1]
        unit, ftype = _text_blank_unit(after)
        # Метка — хвост предложения перед прочерком («Содержание влаги»).
        tail = re.split(r"[.!?]", before)[-1]
        tail = _clean(tail).rstrip(" :=(")
        label = _short_label(" · ".join(x for x in (context, tail) if x) or f"Значение {i + 1}", 90)
        fi = len(fields)
        fields.append({"label": label, "type": ftype, "unit": unit})
        blanks.append({"field_index": fi, "type": ftype, "unit": unit})
    return blanks


def _field_type(label: str) -> str:
    low = label.lower()
    if any(k in low for k in ("вес", "количество", "выход", "скорость", "температур", "влажност", "давлен", "кг", " г", "%")):
        return "number"
    return "text"


def _field_unit(label: str) -> str | None:
    low = label.lower()
    if "кг" in low:
        return "кг"
    if "(г" in low or " г" in low:
        return "г"
    if "%" in low:
        return "%"
    if "об/мин" in low:
        return "об/мин"
    if "минут" in low:
        return "мин"
    return None


def _label_needs_input(label: str) -> bool:
    low = label.lower().strip(" :")
    if "___" in low or "____" in low:
        return True
    return any(k in low for k in (
        # «время окон-я», «время окончание/окончания», «дата нач.» и т.п.
        "время нач", "время окон", "дата нач", "дата окон",
        "количество", "вес", "выход", "отклонение", "другие", "согласование",
    ))


def _signature_role(text: str) -> str | None:
    """Распознаёт ячейку-подпись и роль: ДП (оператор), ДОК (контроль), Склад.
    'Выполнено ДП', 'Исполнитель ДП', 'ИСПЫТАНО: ДП', 'Проверено ДОК',
    'УТВЕРЖДЕНО: ДОК', 'Утверждение (Подпись/Дата)', просто 'Подпись'."""
    low = _clean(text).lower()
    if not low:
        return None
    if re.search(r"\bдок\b", low) or "проверил" in low or "проверено" in low or "утвержд" in low:
        return "signature_qa"
    if re.search(r"\bдп\b", low) or "выполн" in low or "исполнит" in low or "испытан" in low:
        return "signature_operator"
    if "склад" in low or "кладовщик" in low:
        return "signature_warehouse"
    if "подпись" in low:
        return "signature_operator"
    return None


def _is_fill_blank(text: str) -> bool:
    """Ячейка с прочерком для ручного ввода: '_________(г или кг)'."""
    return "___" in (text or "")


def _is_footnote(text: str) -> bool:
    """Сноска/примечание под таблицей: '* …', '** Opadry …', 'Примечание: …'.
    Такие абзацы несут смысл (перерасход Opadry +50%, пересчёт по анализу) и не
    должны ни теряться, ни затирать заголовок следующей секции."""
    s = _clean(text)
    return s.startswith("*") or s.lower().startswith(("примеч", "note", "сноска"))


def _merge_header_rows(raw: list[list[str]]) -> list[list[str]]:
    """Свернуть строки-шапки, растиражированные по вертикали (vMerge: одна
    логическая шапка повторяется в нескольких строках, напр. блистер: «Дата и
    время | Склеивание | Качество склеивания | Размер блистера» в 3 строках +
    уточнение «серия/Срок годности»). Вливаем их в одну, беря самое конкретное
    (нижнее) значение по колонке. Останавливаемся на строке критериев/данных.
    Консервативно: сворачиваем, только если ≥3 колонок — вертикальные дубли,
    иначе обычные двухуровневые шапки (Статус отказа над Блок FE/SS/NFE) не
    трогаем."""
    if len(raw) < 3:
        return raw
    ncols = max((len(r) for r in raw), default=0)
    def cell(r, ci):
        return _clean(r[ci]) if ci < len(r) else ""
    merged = [cell(raw[0], ci) for ci in range(ncols)]
    used = 1
    for ri in range(1, len(raw)):
        row = [cell(raw[ri], ci) for ci in range(ncols)]
        if not any(row):
            break  # пустая строка — это данные
        if re.fullmatch(r"\d+", row[0] or ""):
            break  # строка данных (нумерованная)
        same = sum(1 for ci in range(ncols) if not merged[ci] or not row[ci] or merged[ci] == row[ci])
        diff = ncols - same
        dup_cols = sum(1 for ci in range(ncols) if merged[ci] and row[ci] and merged[ci] == row[ci])
        if same > diff and dup_cols >= 3:
            for ci in range(ncols):
                if row[ci]:
                    merged[ci] = row[ci]  # нижнее = более конкретное
            used += 1
        else:
            break
    if used <= 1:
        return raw
    return [merged] + raw[used:]


def _nested_table(table, prefix: str, fields: list[dict]) -> dict:
    raw = [[_cell_text(c) for c in r.cells] for r in table.rows]
    raw = _merge_header_rows(raw)
    headers = [_clean(h) for h in (raw[0] if raw else [])]
    ncols = max((len(r) for r in raw), default=0)
    # Объединённые ячейки Word python-docx отдаёт повторно по каждому столбцу
    # слияния (напр. «Статус отказа» в 3 ячейках). Схлопываем колонку ТОЛЬКО
    # если она совпадает с предыдущей во ВСЕХ строках (настоящее слияние).
    # Сравнение по одной шапке ломает двухуровневые шапки: «Целостность сита»
    # объединена над «Перед просеиванием | После просеивания» — вторая колонка
    # данных терялась.
    def _col(ri: int, ci: int) -> str:
        r = raw[ri]
        return _clean(r[ci]) if ci < len(r) else ""

    keep: list[int] = []
    for ci in range(ncols):
        h = headers[ci] if ci < len(headers) else ""
        if ci > 0 and h and all(_col(ri, ci) == _col(ri, ci - 1) for ri in range(len(raw))):
            continue
        keep.append(ci)

    # Двухуровневая шапка: в r0 есть объединённая группа (≥2 одинаковых соседних
    # ячейки), а в r1 под ней — разные подзаголовки (напр. «Статус отказа» над
    # «Блок FE | Блок SS | Блок NFE»). Тогда r1 — тоже шапка, а данные с r2; иначе
    # подзаголовки-колонки сами становятся полями, а нумерация строк уезжает.
    header_rows = 1
    if len(raw) > 2:
        ci = 0
        while ci < len(keep):
            cj = ci
            v0 = _col(0, keep[ci])
            while v0 and cj + 1 < len(keep) and _col(0, keep[cj + 1]) == v0:
                cj += 1
            if cj > ci:
                subs = [_col(1, keep[k]) for k in range(ci, cj + 1)]
                if all(subs) and len(set(subs)) == len(subs):
                    header_rows = 2
                    break
            ci = cj + 1

    # Матрица «заголовки колонок × подписи строк»: r0 — разные заголовки колонок
    # (напр. Процесс | Загрузка | … | Разгрузка), а col0 в строках данных — это
    # подписи строк (Время нач. 1, Сделано, Прим.), а НЕ поля. Иначе подписи строк
    # («Время нач. 1») сами становятся полями-временем в колонке-подписи.
    label_col0 = False
    if len(keep) >= 3:
        hdr_row = header_rows - 1
        distinct_headers = len({_col(hdr_row, ci) for ci in keep if _col(hdr_row, ci)})
        if distinct_headers >= 3:
            rowlabel_rows = 0
            for ri2 in range(header_rows, len(raw)):
                c0 = _col(ri2, keep[0])
                inner_empty = sum(1 for ci in keep[1:] if not _col(ri2, ci))
                if c0 and inner_empty >= 2:
                    rowlabel_rows += 1
            if rowlabel_rows >= 2:
                label_col0 = True

    def _sig_label(c: str) -> str | None:
        cl = (c or "").lower()
        if _signature_role(c) and len(c) <= 60 and ("подпись" in cl or "испытан" in cl
                or "утвержд" in cl or "исполнит" in cl or "выполн" in cl):
            return _signature_role(c)
        return None

    out_rows = []
    prev_row_sig = False
    for ri, full_row in enumerate(raw):
        row_num = ri - header_rows + 1
        row = [_clean(full_row[ci]) if ci < len(full_row) else "" for ci in keep]
        nonempty = [c for c in row if c]
        merged_text = len(nonempty) > 1 and len(set(nonempty)) == 1
        row_label = next((c for c in row if c), "")
        # Строка блока подписи (ИСПЫТАНО: ДП / УТВЕРЖДЕНО: ДОК + ДАТА + ВРЕМЯ):
        # все смысловые ячейки — это подпись/дата/время. Разбираем особо: одна
        # подпись на роль + Дата + Время, без дублей и пустых «левых» полей
        # (иначе генерик плодит ×2 подписи и числовые поля из пустых ячеек).
        def _is_sig_or_dt(c: str) -> bool:
            cl = c.lower()
            return bool(_sig_label(c)) or "дата" in cl or "время" in cl
        sig_roles = [_sig_label(c) for c in nonempty if _sig_label(c)]
        row_sig = bool(sig_roles) and all(_is_sig_or_dt(c) for c in nonempty)
        row_role = sig_roles[0] if sig_roles else None
        row_role_name = "Исполнитель ДП" if row_role == "signature_operator" else "Проверено ДОК" if row_role == "signature_qa" else ""
        created_slots: set[str] = set()
        out_cells = []
        for ci in keep:
            text = _clean(full_row[ci]) if ci < len(full_row) else ""
            cell: dict = {"text": text}
            # Заголовок колонки = ближайшая непустая ячейка ВЫШЕ по колонке:
            # при двухуровневой шапке это подзаголовок («Перед просеиванием»),
            # а не объединённый верх («Целостность сита»).
            header = ""
            for hri in range(ri - 1, -1, -1):
                vals = [_col(hri, k) for k in keep]
                filled = [v for v in vals if v]
                if filled and len(set(filled)) == 1 and len(filled) > 1:
                    continue  # строка-инструкция, объединённая на всю ширину
                cand = _col(hri, ci)
                if cand:
                    header = cand
                    break
            if not header:
                header = headers[ci] if ci < len(headers) else ""
            # Метку поля НЕ префиксуем длинным заголовком секции — он и так
            # показан над таблицей. В метке оставляем только строку и колонку.
            label_base = _clean(" · ".join(x for x in (prefix, row_label, header) if x))
            create_input = False
            force_type: str | None = None
            header_sig = _signature_role(header)
            self_sig = _signature_role(text) if text else None
            row_ctx = row_label if (row_label and row_label != text) else f"строка {row_num}"
            tl = text.lower()
            if label_col0 and ci == keep[0]:
                pass  # колонка-подписи матрицы — это метка строки, не поле
            elif row_sig:
                # блок подписи: одна подпись + Дата + Время на роль, без дублей
                if _sig_label(text) and "sig" not in created_slots:
                    create_input = True
                    force_type = _sig_label(text)
                    label_base = _clean(" · ".join(x for x in (prefix, text) if x))
                    created_slots.add("sig")
                elif "дата" in tl and "date" not in created_slots:
                    create_input = True
                    force_type = "text"
                    label_base = _clean(" · ".join(x for x in (prefix, row_role_name, "Дата") if x))
                    created_slots.add("date")
                elif "время" in tl and "time" not in created_slots:
                    create_input = True
                    force_type = "text"
                    label_base = _clean(" · ".join(x for x in (prefix, row_role_name, "Время") if x))
                    created_slots.add("time")
                # прочие/пустые ячейки строки подписи — без поля
            elif prev_row_sig and not nonempty:
                pass  # пустая строка-«продолжение» под блоком подписи — без полей
            elif ri >= header_rows and re.match(r"\s*средн(ий|ее)\b", tl):
                # «Средний = ___» — заполняемое поле среднего (одно на строку),
                # дубли «Средний =» в соседних колонках убираем.
                if "avg" not in created_slots:
                    create_input = True
                    force_type = "number"
                    label_base = _clean(" · ".join(x for x in (row_label, "Средний") if x))
                    created_slots.add("avg")
                else:
                    cell["text"] = ""
            elif ri >= header_rows and header_sig and (merged_text or not text):
                # Колонка-подпись («Подпись», «Выполнено ДП», «Проверено ДОК») —
                # пустые ячейки данных под ней становятся слотом э-подписи.
                create_input = True
                force_type = header_sig
                label_base = _clean(" · ".join(x for x in (row_ctx, header) if x))
            elif ri >= header_rows and self_sig and len(text) <= 60 and ("подпись" in text.lower() or "испытан" in text.lower()
                               or "утвержд" in text.lower() or "исполнит" in text.lower()
                               or "выполн" in text.lower() or re.search(r"\bд[оп]к?\b", text.lower())):
                # Ячейка, которая сама и есть метка подписи («ИСПЫТАНО: ДП»).
                create_input = True
                force_type = self_sig
                label_base = _clean(" · ".join(x for x in (prefix, text) if x))
            elif _is_fill_blank(text):
                # Прочерк под ручной ввод: «_________(г или кг)».
                create_input = True
                clean_txt = _clean(re.sub(r"_+", " ", text))
                label_base = _clean(" · ".join(x for x in (prefix, row_ctx if row_ctx != f"строка {row_num}" else "", header, clean_txt) if x)) or clean_txt
            elif ri >= header_rows and not nonempty and header:
                create_input = True
                label_base = _clean(" · ".join(x for x in (prefix, f"строка {row_num}", header) if x))
            elif ri >= header_rows and not merged_text and not text and nonempty:
                create_input = True
            elif ri >= header_rows and text and not merged_text and _label_needs_input(text):
                create_input = True
                label_base = _clean(" · ".join(x for x in (prefix, text) if x))
            elif len(raw) == 1 and text and _label_needs_input(text):
                create_input = True
                label_base = _clean(" · ".join(x for x in (prefix, text) if x))

            if create_input:
                field_index = len(fields)
                ftype = force_type or _field_type(label_base or text)
                funit = None if force_type else _field_unit(label_base or text)
                fields.append({
                    "label": _short_label(label_base or prefix or f"Поле {field_index + 1}"),
                    "type": ftype,
                    "unit": funit,
                })
                cell["field_index"] = field_index
                cell["type"] = ftype
                if funit:
                    cell["unit"] = funit
            out_cells.append(cell)
        out_rows.append({"cells": out_cells})
        prev_row_sig = row_sig
    return {"rows": out_rows}


def _unique_nested_tables(cells) -> list:
    tables = []
    seen: set[int] = set()
    for cell in cells:
        for nested in cell.tables:
            key = id(nested._tbl)
            if key in seen:
                continue
            seen.add(key)
            tables.append(nested)
    return tables


def _generic_process_table(table, prefix: str) -> dict:
    fields: list[dict] = []
    parsed = _nested_table(table, prefix, fields)
    nested = [
        _nested_table(nt, "Вложенная таблица", fields)
        for nt in _unique_nested_tables(cell for row in table.rows for cell in row.cells)
    ]
    out = {"rows": parsed.get("rows", []), "fields": fields}
    if nested:
        out["tables"] = nested
    if _is_efficiency_calculation(out):
        out = _efficiency_calculation_table(out)
    return out


def _is_efficiency_calculation(parsed: dict) -> bool:
    text = " ".join(
        str(cell.get("text") or "")
        for row in (parsed.get("rows") or [])
        for cell in (row.get("cells") or [])
        if isinstance(cell, dict)
    ).lower()
    return "расчет эффективности" in text or "расчёт эффективности" in text


def _efficiency_calculation_table(parsed: dict) -> dict:
    original_text = " ".join(
        str(cell.get("text") or "")
        for row in (parsed.get("rows") or [])
        for cell in (row.get("cells") or [])
        if isinstance(cell, dict) and cell.get("text")
    )
    return {
        **parsed,
        "process_table_variant": "efficiency_calculation",
        "original_text": original_text,
        "formula": "C1 = 8,651 кг - ((A - 67,425 кг) + (B - 4,928 кг))",
        "standard_metformin_kg": "67,425",
        "standard_sitagliptin_kg": "4,928",
        "standard_mcc_kg": "8,651",
        "fields": [
            {"label": "Метформин · партия / серия", "type": "text"},
            {"label": "Метформин · количественное содержание", "type": "number", "unit": "%"},
            {"label": "Метформин · количество воды", "type": "number", "unit": "%"},
            {"label": "Метформин · фактическое количество (A)", "type": "number", "unit": "кг"},
            {"label": "Ситаглиптин · партия / серия", "type": "text"},
            {"label": "Ситаглиптин · количественное содержание", "type": "number", "unit": "%"},
            {"label": "Ситаглиптин · количество воды", "type": "number", "unit": "%"},
            {"label": "Ситаглиптин · фактическое количество (B)", "type": "number", "unit": "кг"},
            {"label": "Микрокристаллическая целлюлоза · расчетное количество (C1)", "type": "number", "unit": "кг"},
            {"label": "Примечания", "type": "text"},
            {"label": "Рассчитал ДП", "type": "signature_operator"},
            {"label": "Проверил ДОК", "type": "signature_qa"},
        ],
    }


# Коэффициенты «количество на серию из листа распределения» (кг) — на бумаге
# они внутри формулы-дроби (OMML-объект Word), текст которой python-docx не
# отдаёт, поэтому для известных материалов держим их здесь.
_REQ_CALC_COEFF = {
    "EXC-AEROSIL": "0,714",
    "EXC-MCC": "8,651",
    "EXC-SOD": "1,334",
    "EXC-MGST": "1,111",
}
# Стандартный выход гранул на серию (кг) — знаменатель формулы (тоже в OMML).
_REQ_CALC_DIVISOR = "83,379"


def _requirement_calc_table(table, prefix: str, fields: list[dict]) -> dict | None:
    """Расчёт требуемого количества субстанций для опудривания (этапы 11.3/12.3/13.3).

    Формула на бумаге: требуемое = (факт. выход^ × K*) ÷ 83,379; на выброс =
    K* − (взято). Дробь в Word — математический объект, текст не извлекается,
    поэтому разбираем по строкам-парам и строим структуру с 3 полями на материал.
    """
    raw = [[_cell_text(c) for c in r.cells] for r in table.rows]
    flat = " ".join(" ".join(r) for r in raw).lower()
    if "требуемое количество" not in flat or "выбрасыван" not in flat:
        return None

    items: list[dict] = []
    current: dict | None = None
    for r in raw:
        ded: list[str] = []
        for c in r:
            c = _clean(c)
            if c and (not ded or ded[-1] != c):
                ded.append(c)
        if not ded:
            continue
        first = ded[0]
        low = first.lower()
        m = re.search(r"фактическое требуемое количество\s+(.+)$", first, re.IGNORECASE)
        if m:
            name = _clean(m.group(1).rstrip(":"))
            current = {"name": name}
            items.append(current)
            continue
        if "выбрасыван" in low and current is not None:
            rest = " ".join(ded[1:])
            mc = re.search(r"=\s*([\d]+[.,]\d+)\s*\*", rest)
            if mc:
                current["coeff"] = mc.group(1)
            current = None

    if not items:
        return None

    for item in items:
        if not item.get("coeff"):
            alias = _material_alias_key(item["name"])
            item["coeff"] = _REQ_CALC_COEFF.get(alias, "")
        base = item["name"]
        item["required_fi"] = len(fields)
        fields.append({"label": _short_label(f"{base} · требуемое количество"), "type": "number", "unit": "кг"})
        item["taken_fi"] = len(fields)
        fields.append({"label": _short_label(f"{base} · фактически взято"), "type": "number", "unit": "кг"})
        item["discard_fi"] = len(fields)
        fields.append({"label": _short_label(f"{base} · на выброс"), "type": "number", "unit": "кг"})

    return {
        "process_table_variant": "requirement_calculation",
        "divisor": _REQ_CALC_DIVISOR,
        "items": items,
        "rows": [],
    }


def _moisture_loss_table(table, prefix: str, fields: list[dict]) -> dict | None:
    """Лист «Определение влаги и потери при сушке» (этапы 6.6/8.6/10.6).

    Свободная вёрстка Word с разделами a/b/c/d и формулой (a+b+c)×100/d.
    Генерик-парсер размазывает её в десятки дублей — разбираем осмысленно:
    влага и потери по стадиям + итоги, фактический выход, ожидаемый вес и
    результат-процент.
    """
    raw = [[_cell_text(c) for c in r.cells] for r in table.rows]
    flat = " ".join(" ".join(r) for r in raw).lower()
    if "определение влаги" not in flat or "потери" not in flat:
        return None
    if "ожидаемый вес" not in flat and "фактический выход" not in flat:
        return None

    # Названия стадий из строки с «Грануляция/Сушка/…» (без «b.», «Стадия», «Всего»).
    stages: list[str] = []
    for r in raw:
        ded: list[str] = []
        for c in r:
            c = _clean(c)
            if c and (not ded or ded[-1] != c):
                ded.append(c)
        low = " ".join(ded).lower()
        if ("грануляц" in low or "сушк" in low or "смешив" in low or "опудрив" in low) and "потери" not in low:
            for x in ded:
                xl = x.lower().strip(". ")
                if xl in ("a", "b", "c", "d", "стадия", "всего", "образец") or "потери" in xl or "определение" in xl:
                    continue
                if x not in stages:
                    stages.append(x)
            if stages:
                break
    if not stages:
        stages = ["Грануляция", "Сушка"]

    def add(label: str, unit: str | None) -> int:
        fi = len(fields)
        fields.append({"label": _short_label(f"{prefix} · {label}"), "type": "number", "unit": unit})
        return fi

    moisture = [{"stage": s, "fi": add(f"Определение влаги · {s}", "%")} for s in stages]
    moisture_total_fi = add("Определение влаги · Всего (a)", "%")
    losses = [{"stage": s, "fi": add(f"Потери · {s}", "кг")} for s in stages]
    losses_total_fi = add("Потери · Всего (b)", "кг")
    actual_yield_fi = add("Фактический выход (c)", "кг")
    expected_weight_fi = add("Ожидаемый вес (d)", "кг")
    drying_limit_fi = add("Предел отчётности по сушке", "%")

    return {
        "process_table_variant": "moisture_loss",
        "stages": stages,
        "moisture": moisture,
        "moisture_total_fi": moisture_total_fi,
        "losses": losses,
        "losses_total_fi": losses_total_fi,
        "actual_yield_fi": actual_yield_fi,
        "expected_weight_fi": expected_weight_fi,
        "drying_limit_fi": drying_limit_fi,
        "rows": [],
    }


def _yield_calc_table(table, prefix: str, fields: list[dict]) -> dict | None:
    """Лист «Расчёт выхода» (этап 14.7 и т.п.): один столбец из 7 строк —
    Теор. выход (A), Факт. выход (B), Отобранное (C), % выхода=(B+C)/A×100,
    Отклонение (E), Другие (F), Согласование=(B+C+E+F)/A×100. В Word строки —
    `['', метка, '']`, поэтому генерик плодит ×3 поля; разбираем по строкам."""
    raw = [[_cell_text(c) for c in r.cells] for r in table.rows]
    flat = " ".join(" ".join(r) for r in raw).lower()
    if "теоретический выход" not in flat or "согласование" not in flat:
        return None

    def add(label: str, unit: str | None) -> int:
        fi = len(fields)
        fields.append({"label": _short_label(f"{prefix} · {label}"), "type": "number", "unit": unit})
        return fi

    item: dict = {"process_table_variant": "yield_calculation", "rows": []}
    for r in raw:
        label = next((_clean(x) for x in r if _clean(x)), "")
        low = label.lower()
        if not low:
            continue
        if "теоретическ" in low and "выход" in low:
            item["theoretical_fi"] = add("Теоретический выход (A)", "кг")
        elif "согласование" in low:
            item["reconcile_fi"] = add("Согласование = (B+C+E+F)/A×100", "%")
        elif "выход" in low and ("x100" in low or "×100" in low or "b + c" in low or "/ a" in low or "%" in low):
            item["yield_pct_fi"] = add("% выхода = (B+C)/A×100", "%")
        elif "фактическ" in low and "выход" in low:
            item["actual_fi"] = add("Фактический выход (B)", "кг")
        elif "отобранн" in low:
            item["sampled_fi"] = add("Отобранное количество (C)", "кг")
        elif "отклонение" in low:
            item["deviation_fi"] = add("Отклонение (E)", "кг")
        elif "другие" in low:
            item["other_fi"] = add("Другие (F)", "кг")
    if "theoretical_fi" not in item or "reconcile_fi" not in item:
        return None
    return item


def _punch_check_table(table, prefix: str, fields: list[dict]) -> dict | None:
    """Детальная проверка пуансонов (этап 15.1): строки-номера «Верхний/Нижний
    пуансон №: 1..N» + строки «Замечание» (√ соответствует / ✕ нет). Генерик
    плодит призрачные поля из хвостовых пустых ячеек — делаем по одному полю на
    каждый реальный номер пуансона."""
    raw = [[_cell_text(c) for c in r.cells] for r in table.rows]
    flat = " ".join(" ".join(r) for r in raw).lower()
    if "пуансон" not in flat or "замеч" not in flat:
        return None
    ptype = "Верхний пуансон" if "верхний" in flat else ("Нижний пуансон" if "нижний" in flat else "Пуансон")
    items: list[dict] = []
    seen: set[str] = set()
    for row in raw:
        for x in row:
            num = _clean(x)
            if re.fullmatch(r"\d+", num) and num not in seen:
                seen.add(num)
                fi = len(fields)
                fields.append({"label": _short_label(f"{prefix} · {ptype} № {num}"), "type": "text"})
                items.append({"num": num, "fi": fi})
    if not items:
        return None
    items.sort(key=lambda it: int(it["num"]))
    return {"process_table_variant": "punch_check", "ptype": ptype, "punch_items": items, "rows": []}


def _batch_yield_table(table, prefix: str, fields: list[dict]) -> dict | None:
    """Лист «Расчёт выхода серии» на упаковке (этап 9.1): свободная вёрстка
    Word с двумя колонками значений на строку (кг и таблетки) и формулами
    L = ([C+E+F]×100/A) и M = ([C+E+F+G+H+I+J]/B×100). Генерик-парсер
    размазывает объединённые ячейки в десятки дублей; разбираем по строкам:
    константа A (теор. объём в таблетках), измеряемые строки B,C,E…K
    (кг + таблетки), производный итог и два расчётных процента."""
    raw = [[_cell_text(c) for c in r.cells] for r in table.rows]
    flat = " ".join(" ".join(r) for r in raw).lower()
    if "теоретическ" not in flat or "объ" not in flat:
        return None
    if "сверка" not in flat and "% от теоретическ" not in flat:
        return None

    def add(label: str, unit: str | None) -> int:
        fi = len(fields)
        fields.append({"label": _short_label(f"{prefix} · {label}"), "type": "number", "unit": unit})
        return fi

    def clean_label(s: str) -> str:
        s = re.sub(r"_+", " ", s or "")
        s = re.sub(r"\(\s*(кг|г)\s*\)", "", s, flags=re.I)
        s = re.sub(r"\(\s*таблет\w*\s*\)", "", s, flags=re.I)
        s = re.sub(r"^\s*[A-ZА-Я]\s*[.)]\s*", "", _clean(s))  # ведущая буква «K.», «L .»
        # схлопнуть подряд повторённое слово («…оболочкой оболочкой» из merge)
        s = re.sub(r"\b(\w+)(\s+\1\b)+", r"\1", _clean(s), flags=re.I)
        return _clean(s)

    letters = ["B", "C", "E", "F", "G", "H", "I", "J", "K"]
    li = 0
    lines: list[dict] = []
    theoretical_tab: int | None = None
    result_tab_fi: int | None = None
    yield_pct_fi: int | None = None
    reconcile_fi: int | None = None

    for r in raw:
        low = " ".join(r).lower()
        if not low.strip():
            continue
        col0 = _clean(r[0]) if r else ""
        if "теоретическ" in low and "объ" in low:
            digits = re.findall(r"\d[\d\s ]*\d|\d", " ".join(r))
            if digits:
                theoretical_tab = int(re.sub(r"\D", "", digits[-1]) or 0) or None
            continue
        if "% от теоретическ" in low:
            yield_pct_fi = add("% от теоретического выхода ([C+E+F]×100/A)", "%")
            continue
        if "сверка" in low:
            reconcile_fi = add("Сверка ([C+E+F+G+H+I+J]/B×100)", "%")
            continue
        if not col0:  # производный итог: строка без подписи, только «= … (таблеток)»
            result_tab_fi = add("Расчётное количество (таблеток)", "шт")
            continue
        if li < len(letters):  # измеряемая строка: кг + таблетки
            letter = letters[li]
            li += 1
            name = clean_label(col0) or letter
            lines.append({
                "letter": letter,
                "label": name,
                "kg_fi": add(f"{letter}. {name} (кг)", "кг"),
                "tab_fi": add(f"{letter}. {name} (таблеток)", "шт"),
            })

    if not lines or yield_pct_fi is None or reconcile_fi is None:
        return None

    return {
        "process_table_variant": "batch_yield_calculation",
        "rows": [],
        "theoretical_tab": theoretical_tab,
        "lines": lines,
        "result_tab_fi": result_tab_fi,
        "yield_pct_fi": yield_pct_fi,
        "reconcile_fi": reconcile_fi,
    }


def _row_nested_tables(row, prefix: str, fields: list[dict]) -> list[dict]:
    out: list[dict] = []
    for nt in _unique_nested_tables(row.cells):
        special = (
            _requirement_calc_table(nt, prefix, fields)
            or _moisture_loss_table(nt, prefix, fields)
            or _batch_yield_table(nt, prefix, fields)
            or _yield_calc_table(nt, prefix, fields)
            or _punch_check_table(nt, prefix, fields)
        )
        out.append(special if special is not None else _nested_table(nt, prefix, fields))
    return out


def _plain_table(table) -> dict:
    return {
        "rows": [
            {"cells": [{"text": _cell_text(cell)} for cell in row.cells]}
            for row in table.rows
        ],
        "fields": [],
    }


def _room_labels(room_no: str | None) -> list[str]:
    labels: list[str] = []
    for num in re.findall(r"\d{2,3}", room_no or ""):
        label = f"Комн. {num}"
        if label not in labels:
            labels.append(label)
    return labels


def _formula_fingerprint(rows: list[dict]) -> str:
    return "|".join(_norm_key(row.get("group") or row.get("name") or "") for row in rows)


def _formula_score(rows: list[dict]) -> int:
    score = 0
    for row in rows:
        score += len(row.get("name") or row.get("group") or "")
        score += len(row.get("per_tab") or "")
        score += len(row.get("per_series") or "")
    return score


def _process_header(rows: list[list[str]]) -> dict | None:
    """Таблица 'Процесс: X / Комната … / Комната №: N' (ячейки объединённые)."""
    flat = " ".join(_clean(c) for r in rows for c in r)
    low = flat.lower()
    # Настоящая шапка стадии: 'Процесс: …' + 'Комната' + строка 'Дата/Время начала'.
    # Иначе слово «процесс» внутри шагов («в процессе…») ложно детектится.
    if not (re.search(r"процесс\s*:", low) and "комната" in low and ("дата нач" in low or "время нач" in low)):
        return None
    proc = room = room_no = ""
    m = re.search(r"процесс\s*:?\s*(.+?)\s*(?:комната|дата|$)", flat, re.IGNORECASE)
    if m:
        proc = _clean(m.group(1))
    m = re.search(r"комната\s*№\s*:?\s*([^|]*?)(?:\s*(?:дата|время|предыд)|$)", flat, re.IGNORECASE)
    if m:
        room_no = _clean(m.group(1))
    m = re.search(r"комната\s*:?\s*(?!№)(.+?)\s*(?:комната\s*№|дата|$)", flat, re.IGNORECASE)
    if m:
        room = _clean(m.group(1))
    return {"process": proc, "room": room, "room_no": room_no}


def _expand_column_time(tbl: dict) -> list[dict] | None:
    """Колоночный стиль времени: шапка «Время начала | Время окончания | данные»
    с пустыми ячейками-полями под ними (напр. «Отбор проб»: начала|окончания|
    количество). Разворачиваем в отдельные блоки, чтобы _split_time_tables смог
    их разнести: [Время начала] · [данные (количество)] · [Время окончания].
    Возвращает None, если это не колоночный стиль или строк данных несколько."""
    rows = tbl.get("rows") or []
    col_start = col_end = None
    col_hdr: dict[int, str] = {}
    for row in rows:
        for ci, cell in enumerate(row.get("cells") or []):
            if "field_index" in cell:
                continue
            low = _clean(str(cell.get("text") or "")).lower()
            if "время нач" in low or "дата нач" in low:
                col_start = ci; col_hdr[ci] = _clean(str(cell.get("text")))
            elif "время окон" in low or "дата окон" in low:
                col_end = ci; col_hdr[ci] = _clean(str(cell.get("text")))
    if col_start is None and col_end is None:
        return None

    def fields_in_col(ci: int | None) -> list[dict]:
        if ci is None:
            return []
        out = []
        for row in rows:
            cells = row.get("cells") or []
            if ci < len(cells) and "field_index" in cells[ci]:
                out.append(cells[ci])
        return out

    starts = fields_in_col(col_start)
    ends = fields_in_col(col_end)
    if not starts and not ends:
        return None
    if len(starts) > 1 or len(ends) > 1:
        return None  # несколько строк данных — не трогаем (риск потери)

    time_cols = {c for c in (col_start, col_end) if c is not None}
    blocks: list[dict] = []
    if starts:
        sc = dict(starts[0]); sc["text"] = col_hdr.get(col_start, "Время начала")
        blocks.append({"rows": [{"cells": [sc]}]})
    mid_rows = [{"cells": [c for ci, c in enumerate(row.get("cells") or []) if ci not in time_cols]}
                for row in rows]
    if any("field_index" in c for r in mid_rows for c in r["cells"]):
        blocks.append({"rows": mid_rows})
    if ends:
        ec = dict(ends[0]); ec["text"] = col_hdr.get(col_end, "Время окончания")
        blocks.append({"rows": [{"cells": [ec]}]})
    return blocks if len(blocks) > 1 else None


def _split_time_tables(tables: list[dict]) -> list[dict]:
    """Время начала — ДО выполнения, окончания — ПОСЛЕ.

    По процессу: отметил время начала → выполнил/взвесил → отметил окончание.
    В Word «Время начала | окончания» идут одной строкой (часто после данных),
    поэтому разносим: начало ставим перед данными блока, окончание — после.
    Поддержаны и один блок (2.5), и несколько (14.5 — три взвешивания, у
    каждого своё время)."""
    # Сначала разворачиваем колоночный стиль времени в отдельные блоки.
    expanded: list[dict] = []
    for tbl in tables:
        cols = _expand_column_time(tbl)
        expanded.extend(cols if cols else [tbl])
    tables = expanded

    def kind(cell: dict) -> str | None:
        low = _clean(str(cell.get("text") or "")).lower()
        if "field_index" not in cell:
            return None
        if "время нач" in low or "дата нач" in low:
            return "start"
        if "время окон" in low or "дата окон" in low:
            return "end"
        return None

    def split_cells(tbl: dict) -> tuple[list[dict], list[dict], bool]:
        starts: list[dict] = []
        ends: list[dict] = []
        other = False
        for row in tbl.get("rows") or []:
            for cell in row.get("cells") or []:
                k = kind(cell)
                if k == "start":
                    starts.append(cell)
                elif k == "end":
                    ends.append(cell)
                elif "field_index" in cell:
                    other = True  # настоящий ввод (не время) — блок не «чисто временной»
                elif cell.get("text"):
                    # голая метка «Дата:» / «Время:» (без поля) не должна мешать
                    # разносу начала/окончания по блокам — иначе оба времени
                    # остаются вместе и «заполняются сразу».
                    lbl = _clean(str(cell.get("text"))).lower().strip(" :")
                    if lbl not in ("дата", "время", "дата и время", "дата/время"):
                        other = True
        return starts, ends, other

    time_only = sum(1 for t in tables if (lambda r: r[2] is False and (r[0] or r[1]))(split_cells(t)))
    if time_only == 0:
        return tables

    if time_only == 1:
        # Один блок времени на весь шаг: начало — в самое начало, окончание — в конец.
        head: list[dict] = []
        middle: list[dict] = []
        tail: list[dict] = []
        for tbl in tables:
            starts, ends, other = split_cells(tbl)
            if (starts or ends) and not other:
                if starts:
                    head.append({"rows": [{"cells": starts}]})
                if ends:
                    tail.append({"rows": [{"cells": ends}]})
            else:
                middle.append(tbl)
        return head + middle + tail

    # Несколько блоков (напр. 14.5): для каждого — начало перед его данными, окончание после.
    result: list[dict] = []
    last_data_idx: int | None = None  # позиция блока данных, к которому привязываем время
    for tbl in tables:
        starts, ends, other = split_cells(tbl)
        if (starts or ends) and not other:
            if starts:
                pos = last_data_idx if last_data_idx is not None else len(result)
                result.insert(pos, {"rows": [{"cells": starts}]})
                if last_data_idx is not None:
                    last_data_idx += 1
            if ends:
                result.append({"rows": [{"cells": ends}]})
        else:
            result.append(tbl)
            last_data_idx = len(result) - 1
    return result


def _steps(table) -> tuple[list[dict], list[dict]]:
    """Таблица '№ | ТЕХНОЛОГИЧЕСКИЕ ЭТАПЫ | Подпись' → шаги."""
    steps = []
    fields: list[dict] = []
    for row in table.rows:
        raw_cells = [_cell_text(c) for c in row.cells]
        # Объединённые ячейки Word python-docx отдаёт повторно по каждому столбцу
        # слияния: строка этапа выглядит как ['3.1','3.1','Установите…','Установите…'].
        # Схлопываем подряд идущие дубли, иначе text=cells[1] берёт повтор номера,
        # а сама инструкция (cells[2]) теряется → в eBMR виден только «3.1».
        cells: list[str] = []
        for c in raw_cells:
            if c and (not cells or cells[-1] != c):
                cells.append(c)
        if not cells:
            continue
        head = " ".join(cells).lower()
        if "технолог" in head and "этап" in head:
            continue
        no = cells[0]
        text = cells[1] if len(cells) > 1 else ""
        if re.match(r"^\d+(\.\d+)*$", no):
            tables = _row_nested_tables(row, f"Этап {no}", fields)
            # Продолжение таблицы того же шага на новой странице: текст либо
            # пуст, либо состоит из одной бумажной строки подписи («Выполнил
            # ДП…») — сравниваем ПОСЛЕ очистки, иначе плодится дубль-шаг.
            stripped = _strip_signature_text(text)
            if len(stripped) <= 2 and tables and steps and steps[-1].get("no") == no:
                steps[-1].setdefault("tables", []).extend(tables)
                continue
            if len(stripped) <= 2 and not tables:
                continue
            step = {"no": no, "text": stripped or "Контрольная таблица"}
            text_blanks = _extract_text_blanks(stripped, fields, f"Этап {no}")
            if text_blanks:
                step["text_blanks"] = text_blanks
            if tables:
                step["tables"] = tables
            step["dp_field_index"] = len(fields)
            fields.append({"label": f"Этап {no} · Выполнено ДП", "type": "signature_operator"})
            step["dok_field_index"] = len(fields)
            fields.append({"label": f"Этап {no} · Проверено ДОК", "type": "signature_qa"})
            steps.append(step)
    for step in steps:
        if step.get("tables"):
            step["tables"] = _split_time_tables(step["tables"])
    return steps, fields


def _line_clearance_steps_from_rows(rows: list[list[str]]) -> tuple[list[dict], list[dict]]:
    """Line clearance sometimes has typoed headers and empty merged cells.
    Treat it as a GMP checklist: each row is signed by DP, then verified by DOK.
    """
    steps: list[dict] = []
    fields: list[dict] = []
    for row in rows:
        cells = [_clean(c) for c in row]
        if not any(cells):
            continue
        row_text = " ".join(cells).lower()
        if "технолог" in row_text and ("этап" in row_text or "подпись" in row_text):
            continue
        no = cells[0] if cells else ""
        text = cells[1] if len(cells) > 1 else " ".join(cells[1:])
        if not text or text.lower() in {"технологические этапы", "технологоческие этапы"}:
            continue
        if not no or not re.search(r"\d", no):
            no = str(len(steps) + 1)
        no = no.strip(".")
        step = {"no": no, "text": text, "dp_field_index": len(fields)}
        fields.append({"label": f"Этап {no} · Выполнено ДП", "type": "signature_operator"})
        step["dok_field_index"] = len(fields)
        fields.append({"label": f"Этап {no} · Проверено ДОК", "type": "signature_qa"})
        steps.append(step)
    if steps:
        fields.append({"label": "Итоговое утверждение line clearance · ДОК", "type": "signature_qa"})
    return steps, fields


def _looks_like_line_clearance_checklist(title: str, flat: str, steps: list[dict]) -> bool:
    text = " ".join([title, flat, " ".join(str(step.get("text") or "") for step in steps)]).lower()
    markers = (
        "контрольная таблица очист",
        "весь персонал",
        "ничего не осталось от предыдущей серии",
        "стикер «зелёный»",
        "стикер \"зелёный\"",
        "журналы для конкретной машины",
    )
    return any(marker in text for marker in markers)


def _line_clearance_extra(extra: dict, stage: str, title: str = "", flat: str = "") -> dict:
    steps = extra.get("steps") or []
    if not stage.startswith("line_clearance") or not _looks_like_line_clearance_checklist(title, flat, steps):
        return extra
    fields = list(extra.get("fields") or [])
    approval_index = next(
        (
            idx for idx, field in enumerate(fields)
            if "итоговое утверждение" in str((field or {}).get("label") or "").lower()
            and "line clearance" in str((field or {}).get("label") or "").lower()
        ),
        None,
    )
    if approval_index is None:
        approval_index = len(fields)
        fields.append({"label": "Итоговое утверждение line clearance · ДОК", "type": "signature_qa"})
    return {
        **extra,
        "fields": fields,
        "line_clearance_checklist": True,
        "approval_field_index": approval_index,
    }


def _env_params(rows: list[list[str]]) -> list[dict]:
    params = []
    for r in rows[1:]:
        name = _clean(r[0]) if r else ""
        if name and "параметр" not in name.lower():
            params.append({"name": name, "limit": ""})
    return params


def _equipment(rows: list[list[str]]) -> list[dict]:
    out = []
    headers = [_clean(c).lower() for c in (rows[0] if rows else [])]
    has_brand = any("марка" in h for h in headers)
    has_combined_model_serial = any("модель/сер" in h or ("модель" in h and "сер" in h) for h in headers)
    for r in rows[1:]:
        cells = [_clean(c) for c in (r + [""] * 7)[:7]]
        name = _clean(cells[1]) or _clean(cells[0])
        if not name or name.lower() in ("№", "no"):
            continue
        if has_brand and len(cells) >= 7:
            out.append({
                "name": name,
                "model": cells[2],
                "brand": cells[3],
                "serial": cells[4],
                "sop": cells[5],
                "calib": cells[6],
            })
            continue
        # Старый формат бумажного BMR: "Модель/сер.№" в одной колонке,
        # затем СОП и дата поверки. Не раскладываем СОП в серийный номер.
        if has_combined_model_serial or len([c for c in cells if c]) <= 5:
            out.append({
                "name": name,
                "model": cells[2],
                "brand": "",
                "serial": "",
                "sop": cells[3],
                "calib": cells[4],
            })
            continue
        out.append({
            "name": name,
            "model": cells[2],
            "brand": cells[3],
            "serial": cells[4],
            "sop": cells[5],
            "calib": cells[6],
        })
    return out


def _formula(rows: list[list[str]]) -> list[dict]:
    out = []
    seen_materials: set[tuple[str, str]] = set()
    for r in rows[1:]:
        cells = (r + [""] * 6)[:6]
        name = _clean(cells[2])
        if not name:
            continue
        # строка-группа: все смысловые ячейки одинаковы
        vals = [c for c in cells if c]
        if len(set(vals)) == 1:
            out.append({"group": name})
            continue
        item_key = (_norm_key(name), _clean(cells[5]))
        # В Word одна и та же позиция часто идёт второй строкой без количества
        # для ручного вписывания второй серии. В eBMR это даёт дубли при FEFO,
        # поэтому пустой дубль не переносим как отдельный ингредиент.
        if not _clean(cells[5]) and any(k[0] == item_key[0] for k in seen_materials):
            continue
        seen_materials.add(item_key)
        out.append({"name": name, "spec": _clean(cells[3]),
                    "per_tab": _clean(cells[4]), "per_series": _clean(cells[5])})
    # схлопываем подряд идущие дубли (одна и та же строка повторяется в шаблоне)
    dedup = []
    for it in out:
        if dedup and dedup[-1] == it:
            continue
        dedup.append(it)
    return dedup


def _formula_lookup(rows: list[dict]) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for row in rows:
        name = row.get("name")
        if name:
            out.setdefault(_norm_key(name), row)
    return out


def _find_formula(name: str, formula_by_name: dict[str, dict]) -> dict:
    key = _norm_key(name)
    if key in formula_by_name:
        return formula_by_name[key]
    key_tokens = set(key.split())
    best: tuple[float, dict] = (0.0, {})
    for formula_key, row in formula_by_name.items():
        formula_tokens = set(formula_key.split())
        if key_tokens and formula_tokens and (
            key_tokens.issubset(formula_tokens) or formula_tokens.issubset(key_tokens)
        ):
            return row
        score = SequenceMatcher(None, key, formula_key).ratio()
        if score > best[0]:
            best = (score, row)
    return best[1] if best[0] >= 0.72 else {}


def _distribution(rows: list[list[str]], formula_by_name: dict[str, dict]) -> list[dict]:
    groups: list[dict] = []
    current: dict | None = None
    seen_in_group: set[str] = set()

    for r in rows[1:]:
        cells = (r + [""] * 8)[:8]
        nonempty = [_clean(c) for c in cells if _clean(c)]
        if not nonempty:
            continue
        flat = " ".join(nonempty).lower()
        if "аналит" in flat and ("серия" in flat or "наименован" in flat):
            continue
        if "подпись" in flat or "выдано" in flat or "проверено" in flat:
            continue

        # Групповая строка в исходном BMR заполнена одним и тем же текстом во
        # всех объединённых ячейках.
        if len(nonempty) >= 3 and len(set(nonempty)) == 1 and not _clean(cells[2]):
            title = nonempty[0]
            current = {"title": title, "items": []}
            groups.append(current)
            seen_in_group = set()
            continue
        if len(nonempty) >= 3 and len(set(nonempty)) == 1 and _clean(cells[2]) == nonempty[0]:
            title = nonempty[0]
            current = {"title": title, "items": []}
            groups.append(current)
            seen_in_group = set()
            continue

        name = _clean(cells[2])
        if not name or "наименован" in name.lower():
            continue
        norm = _norm_key(name)
        if norm in seen_in_group:
            continue
        seen_in_group.add(norm)
        formula = _find_formula(name, formula_by_name)
        qty = _clean(cells[3]).replace("_", "").strip() or formula.get("per_series") or ""
        spec = formula.get("spec") or ""
        if current is None:
            current = {"title": "Материалы", "items": []}
            groups.append(current)
        current["items"].append({"name": name.replace("*", "").strip(), "spec": spec, "qty": qty})

    return [g for g in groups if g.get("items")]


def _dedupe_labels(fields: list[dict]) -> None:
    """Делает метки полей уникальными в пределах секции. Повторяющиеся блоки
    колонок (напр. несколько тар в ряд) или одинаковые строки-проверки дают
    совпадающие метки («строка 2 · Общий вес» ×4) — нумеруем повторы, чтобы в
    обзоре полей не было визуального дублирования."""
    seen: dict[str, int] = {}
    for f in fields:
        label = f.get("label") or ""
        n = seen.get(label, 0) + 1
        seen[label] = n
        if n > 1:
            f["label"] = f"{label} ({n})"


def _enrich_material_codes(db, sections: list[dict]) -> None:
    materials = db.query(Material).all()
    by_code = {m.code: m for m in materials}
    by_name = {_norm_key(m.name): m for m in materials}

    def resolve(name: str) -> str:
        alias = _material_alias_key(name)
        if alias and alias in by_code:
            return alias
        key = _norm_key(name)
        if key in by_name:
            return by_name[key].code
        key_tokens = set(key.split())
        for material_key, material in by_name.items():
            material_tokens = set(material_key.split())
            if key_tokens and material_tokens and (
                key_tokens.issubset(material_tokens) or material_tokens.issubset(key_tokens)
            ):
                return material.code
        return ""

    for section in sections:
        cfg = section.get("config") or {}
        if cfg.get("kind") == "production_formula":
            for row in cfg.get("rows", []):
                if row.get("name"):
                    row["material_code"] = resolve(row.get("name") or "")
            continue
        if cfg.get("kind") != "distribution_list":
            continue
        field_base = 0
        for group in cfg.get("groups", []):
            for item in group.get("items", []):
                item["field_base"] = field_base
                item["material_code"] = resolve(item.get("name") or "")
                field_base += 6


def build_sections(doc: Document) -> list[dict]:
    sections: list[dict] = []
    ordinal = 0
    cur_stage = "identity"
    cur_room = None
    cur_rooms: list[str] = []
    pending_title = ""
    last_env_heading = ""
    note_buffer: list[str] = []
    formula_by_name: dict[str, dict] = {}
    seen_formula_sections: dict[str, int] = {}
    seen_formula_scores: dict[str, int] = {}
    seen_distribution_fingerprints: set[str] = set()

    def add(stype: str, title: str, kind: str, extra: dict):
        nonlocal ordinal
        cfg = {"kind": kind, "stage": cur_stage, "room": cur_room, **extra}
        if cur_rooms and "rooms" not in cfg:
            cfg["rooms"] = cur_rooms
        if cur_room and not cur_rooms and "room_assignment_required" not in cfg:
            cfg["room_assignment_required"] = True
            cfg["room_source_text"] = cur_room
        # Прочерки «впиши номер комнаты от руки» в заголовках стадий — мусор:
        # комната в eBMR назначается отдельным полем.
        title = _clean(re.sub(r"\s*_{2,}\s*", " ", title or ""))
        sections.append({"ordinal": ordinal, "section_type": stype, "title": _section_title(title, "Контрольная таблица")[:255], "config": cfg})
        # Висячие сноски, встреченные до этой секции, прикрепляем к ней.
        if note_buffer:
            notes = sections[-1]["config"].setdefault("notes", [])
            for n in note_buffer:
                if n not in notes:
                    notes.append(n)
            note_buffer.clear()
        ordinal += 1

    _DUP_IGNORE = ("stage", "room", "rooms", "room_assignment_required", "room_source_text", "notes")

    def dup(stype, kind, extra) -> bool:
        if not sections:
            return False
        last = sections[-1]
        return last["section_type"] == stype and last["config"].get("kind") == kind and \
            {k: v for k, v in last["config"].items() if k not in _DUP_IGNORE} == {"kind": kind, **extra}

    for typ, val in _iter_blocks(doc):
        if typ == "p":
            if _is_footnote(val):
                # Сноску не делаем заголовком — привязываем к уже созданной
                # секции (она идёт под своей таблицей) либо буферизуем.
                if sections:
                    notes = sections[-1]["config"].setdefault("notes", [])
                    if val not in notes:
                        notes.append(val)
                elif val not in note_buffer:
                    note_buffer.append(val)
            else:
                pending_title = val
                # Заголовок отчёта о среде («ОТЧЁТ ОБ УСЛОВИЯХ СРЕДЫ ДЛЯ КОМНАТЫ …»)
                # часто перекрывается следующим под-абзацем — запоминаем отдельно.
                low = val.lower()
                if "сред" in low and ("комнат" in low or "област" in low or "отч" in low):
                    last_env_heading = val
            continue
        rows = _rows(val)
        if not rows:
            continue
        flat = " ".join(" ".join(r) for r in rows).lower()

        ph = _process_header(rows)
        if ph:
            room_labels = _room_labels(ph["room_no"])
            cur_rooms = room_labels
            cur_room = " / ".join(room_labels) if room_labels else (ph["room"] or cur_room)
            proc = ph["process"]
            room_slug = re.sub(r"[^a-zа-я0-9]+", "_", (ph["room"] or ph["room_no"]).lower()).strip("_")[:24]
            proc_slug = re.sub(r"[^a-zа-я0-9]+", "_", proc.lower()).strip("_")[:32]
            if not proc_slug:
                proc_slug = "process"
            if "очистка" in proc.lower():
                # line clearance уникален по комнате (иначе все слипаются)
                cur_stage = f"line_clearance_{room_slug or proc_slug}"
            elif "взвеш" in proc.lower():
                cur_stage = "weighing"
            else:
                cur_stage = f"{proc_slug}_{room_slug}" if room_slug else proc_slug
            header_extra = {"room_no": ph["room_no"], "rooms": room_labels, "process": proc, "stage_title": proc, "fields": _process_fields()}
            if not room_labels:
                header_extra["room_assignment_required"] = True
                header_extra["room_source_text"] = ph["room"] or ph["room_no"] or ""
            add("stage", f"Процесс: {proc} · {ph['room']} {ph['room_no']}".strip(), "process_header", header_extra)
            pending_title = ""
            continue

        # Справочные таблицы шапки BMR. Они нужны руководителю/ДОК как часть
        # master-copy, но не должны становиться заданиями оператора на планшете.
        if "утверждение" in flat and "департамент производства" in flat:
            extra = {**_plain_table(val), "operator_visible": False}
            add("reference_table", "Матрица утверждения BMR", "reference_table", extra)
            pending_title = ""
            continue
        if "ф.и.о" in flat and "инициалы" in flat and "департамент" in flat:
            extra = {**_plain_table(val), "operator_visible": False}
            add("reference_table", "Журнал подписей BMR", "reference_table", extra)
            pending_title = ""
            continue
        if "version" in flat and "reason" in flat:
            extra = {**_plain_table(val), "operator_visible": False}
            add("reference_table", "История изменений BMR", "reference_table", extra)
            pending_title = ""
            continue

        # производственная формула
        if "состав" in flat and ("кол-во" in flat or "количество на серию" in flat):
            formula_rows = _formula(rows)
            if formula_rows and not formula_by_name:
                formula_by_name = _formula_lookup(formula_rows)
            extra = {"rows": formula_rows, "fields": []}
            fingerprint = _formula_fingerprint(formula_rows)
            score = _formula_score(formula_rows)
            if fingerprint in seen_formula_sections:
                section = sections[seen_formula_sections[fingerprint]]
                if score > seen_formula_scores.get(fingerprint, 0):
                    section["config"]["rows"] = formula_rows
                    seen_formula_scores[fingerprint] = score
            elif not dup("production_formula", "production_formula", extra):
                seen_formula_sections[fingerprint] = len(sections)
                seen_formula_scores[fingerprint] = score
                add("production_formula", pending_title or "Производственная формула", "production_formula", extra)
            pending_title = ""
            continue

        # лист расположения / распределения
        if "лист располож" in pending_title.lower() or ("наименован" in flat and "серия субстанц" in flat):
            groups = _distribution(rows, formula_by_name)
            fingerprint = repr(groups)
            if fingerprint in seen_distribution_fingerprints:
                pending_title = ""
                continue
            seen_distribution_fingerprints.add(fingerprint)
            extra = {"stage": "weighing", "groups": groups, "fields": _distribution_fields(groups)}
            if not dup("distribution_list", "distribution_list", extra):
                add("distribution_list", pending_title or "Лист расположения", "distribution_list", extra)
            pending_title = ""
            continue

        # окружающая среда: настоящий отчёт о среде имеет В ШАПКЕ и «Начало», и
        # «Окончание», и колонку «Наблюдение» (Температура/Влажность × старт/конец
        # × Наблюдение/ДП/ДОК). Раньше детект ловил любое «начало»/«окончание» и
        # утаскивал сюда внутрипроцессный контроль «В процессе (Начало)» — таблицу
        # таблеток с 16 параметрами. Требуем все три признака.
        if "параметр" in flat and "начало" in flat and "окончание" in flat and "наблюдение" in flat:
            extra = {"params": _env_params(rows), "fields": _sig_fields()}
            env_title = last_env_heading or pending_title or "Отчёт об условиях окружающей среды"
            if not dup("environment", "environment", extra):
                add("environment", env_title, "environment", extra)
            pending_title = ""
            last_env_heading = ""
            continue

        # оборудование
        if (
            ("оборудован" in flat or "название оборудования" in flat)
            and ("модель" in flat or "модел" in flat or "поверк" in flat or "серия" in flat or "марка" in flat)
        ):
            extra = {"rows": _equipment(rows), "fields": []}
            if not dup("equipment", "equipment", extra):
                add("equipment", pending_title or "Оборудование / приборы", "equipment", extra)
            pending_title = ""
            continue

        # технологические этапы (шаги процесса) с подписями.
        # Таблицы материалов («Использованные упаковочные материалы») тоже
        # имеют нумерованные строки — это НЕ шаги с подписями ДП/ДОК на каждую
        # строку, пропускаем их в универсальную process_table ниже.
        is_material_table = ("упаковочн" in f"{pending_title} {flat}".lower()
                             and "материал" in f"{pending_title} {flat}".lower())
        st, st_fields = ([], []) if is_material_table else _steps(val)
        if st and ("технолог" in flat or "этап" in flat or len(st) >= 1):
            extra = _line_clearance_extra({"steps": st, "fields": st_fields}, cur_stage, pending_title, flat)
            add("checklist", pending_title or "Технологические этапы", "checklist", extra)
            pending_title = ""
            continue

        # line clearance in paper BMR is a checklist with DP execution and DOK
        # verification per row. Do not import it as a free-form process table.
        if cur_stage.startswith("line_clearance") and ("выполнено дп" in flat or "подпись" in flat):
            lc_steps, lc_fields = _line_clearance_steps_from_rows(rows)
            if lc_steps:
                add(
                    "checklist",
                    pending_title or "Контрольная таблица очистки линии",
                    "checklist",
                    _line_clearance_extra({
                        "steps": lc_steps,
                        "fields": lc_fields,
                    }, cur_stage, pending_title or "Контрольная таблица очистки линии", flat),
                )
                pending_title = ""
                continue

        # шапка идентификации
        if "код продукта" in flat:
            kv = {}
            for r in rows:
                if len(r) >= 3 and r[0]:
                    kv[r[0]] = _clean(r[2])
            add("product_header", "Шапка ЗПС / идентификация", "product_header", {"fields": [], "identity": kv})
            pending_title = ""
            continue

        # Бумажное «утверждено ДОК» теперь отражается в каждой строке line
        # clearance как e-подпись «Проверено ДОК», отдельная секция не нужна.
        if "утверждено док" in flat:
            pending_title = ""
            continue

        # Самостоятельные таблицы контроля/расчётов без шапки «Технологические
        # этапы»: металлоискатель, параметры прессования, IPC, блистеровка,
        # упаковочные материалы, итог выхода, финальные подписи и т.п.
        meaningful_source = f"{pending_title.lower()} {flat}"
        meaningful = any(word in meaningful_source for word in (
            "металлоиск", "стандартные параметры", "в процессе", "внешний вид",
            "склеивание", "упаковочные материалы", "итого", "фактический выход",
            "рассмотрено", "подпись", "составлено", "этикетка статуса", "распечатка",
            "расчет эффективности", "параметр испытаний", "признаваемые критерии",
            "этапы процесса", "наименование документа", "статус", "пуансон",
            "матрица", "отбор проб",
        ))
        if meaningful:
            title = _section_title(pending_title, "Контрольная таблица")
            # Префикс полей пустой: длинный текст-инструкция уже хранится в title
            # секции и показывается над таблицей; дублировать его в каждую метку
            # поля не нужно (иначе «<вся инструкция> · строка N · Дата»).
            parsed = _generic_process_table(val, "")
            if parsed["rows"]:
                extra = parsed
                if "расчет эффективности" in meaningful_source and cur_stage == "identity":
                    extra = {**parsed, "stage": "weighing", "room": "Комн. 39", "rooms": ["Комн. 39"]}
                add("process_table", title, "process_table", extra)
            pending_title = ""
            continue

        pending_title = ""
    return sections


def main(path: str, code: str, name: str) -> None:
    data = open(path, "rb").read()
    doc = Document(BytesIO(data))
    sections = build_sections(doc)

    db = SessionLocal()
    try:
        author = db.query(User).filter(User.username == "head_qc").first() or db.query(User).first()
        product = db.query(Product).filter(Product.code == code).first()
        if not product:
            product = Product(code=code, name=name, dosage_form="Таблетки, покрытые оболочкой", market_code="UZ", market_name="Узбекистан")
            db.add(product)
            db.flush()
        for s in sections:
            _dedupe_labels(s["config"].get("fields") or [])
        _enrich_material_codes(db, sections)
        # новая версия (draft)
        from sqlalchemy import func as f
        ver = (db.query(f.max(BmrTemplate.version)).filter(BmrTemplate.product_id == product.id).scalar() or 0) + 1
        tpl = BmrTemplate(product_id=product.id, title=f"ЗПС {name} — импорт из docx", version=ver,
                          status="draft", created_by=author.id,
                          notes="Импортировано из .docx (фабричный формат). ЧЕРНОВИК — проверьте все стадии/параметры/оборудование перед утверждением (СОП-11).")
        db.add(tpl)
        db.flush()
        for s in sections:
            db.add(BmrSection(template_id=tpl.id, ordinal=s["ordinal"], section_type=s["section_type"],
                              title=s["title"], config=s["config"]))
        db.commit()
        stages = []
        for s in sections:
            st = s["config"].get("stage")
            if st and st not in stages:
                stages.append(st)
        print(f"OK template id={tpl.id} v{ver} status=draft product={code} sections={len(sections)}")
        print("stages:", ", ".join(stages))
    finally:
        db.close()


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else sys.argv[2])
