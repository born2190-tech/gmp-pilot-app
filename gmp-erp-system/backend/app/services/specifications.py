"""Сервис справочника спецификаций (НД) на материалы.

Спецификация — это набор показателей качества с нормами и методами для
входного контроля. Аналитический лист ОКК подставляет параметры выбранной
спецификации. Резолвер сопоставляет спецификацию с партией по `material_id`
либо по ключевым словам/наименованию.
"""
from __future__ import annotations

import re
import tempfile
from io import BytesIO
from uuid import UUID

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.api.deps import CurrentUser
from app.models.inventory import Lot
from app.models.master_data import Material
from app.models.quality import MaterialSpecification, SpecificationParameter
from app.schemas.quality import MaterialSpecificationInput
from app.services.audit import write_audit
from app.services.permissions import require_permission


_ND_RE = re.compile(r"\b(?:НД|ND|SPC|СПЦ|ФСП|FSP)[-_/.\w()А-Яа-я]*\d[-_/.\w()А-Яа-я]*", re.IGNORECASE)
_REV_RE = re.compile(r"\b(?:ревизия|revision|редакция)\b\s*[№#:]?\s*([A-Za-zА-Яа-я0-9./-]+)", re.IGNORECASE)
_UNIT_RE = re.compile(r"(%|мг|mg|г|g|кг|kg|мл|ml|КОЕ/г|cfu/g|ppm)\b", re.IGNORECASE)
_HEADER_NOISE_RE = re.compile(
    r"(адрес|телефон|mail|e-mail|ул\.|улица|район|область|республика|завод|novugen|"
    r"спецификация$|ссылка\s+на\s+нд|таблица|страница|утвержд|разработ|соглас)",
    re.IGNORECASE,
)
_CHEM_FORMULA_RE = re.compile(r"^[A-ZА-Я]?\d*[A-Z][A-Za-z]?\d+[A-Za-zА-Яа-я0-9().,\-\s]+$")


def _clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\x00", " ")).strip()


def _filename_stem(filename: str) -> str:
    return _clean(filename.rsplit("\\", 1)[-1].rsplit("/", 1)[-1].rsplit(".", 1)[0])


def _normalise_code(value: str) -> str:
    value = re.sub(r"[\s_]+", "-", _clean(value))
    value = re.sub(r"[^0-9A-Za-zА-Яа-я./_-]+", "-", value)
    value = re.sub(r"-{2,}", "-", value).strip("-_ .")
    return value[:128]


def _code_from_filename(filename: str) -> str | None:
    stem = _filename_stem(filename)
    parts = [p for p in re.split(r"[_\s]+", stem) if p]
    useful: list[str] = []
    for part in parts:
        lower = part.lower()
        if any(stop in lower for stop in ("специфика", "specification", "ситаглип", "дапаг", "тиг", "материал")):
            break
        cleaned = re.sub(r"[^0-9A-Za-zА-Яа-я-]+", "", part)
        if cleaned:
            useful.append(cleaned)
    if useful and any(re.search(r"\d", p) for p in useful) and any(re.match(r"(?i)^(spc|nd|фсп|спц|соп)", p) for p in useful):
        return _normalise_code("-".join(useful))
    return None


def _material_from_filename(filename: str) -> str | None:
    stem = _filename_stem(filename)
    stem = re.sub(r"^\W*\d+[\W_]*", " ", stem)
    stem = re.sub(r"[_\-()]+", " ", stem)
    stop_words = {"spc", "nd", "фсп", "спц", "соп", "spec", "sub", "суб", "specification", "спецификация"}
    stem = " ".join(part for part in stem.split() if not part.isdigit() and part.lower() not in stop_words)
    stem = _clean(stem)
    if 3 <= len(stem) <= 120 and not _HEADER_NOISE_RE.search(stem):
        return stem
    return None


def _has_cyrillic(value: str | None) -> bool:
    return bool(value and re.search(r"[А-Яа-я]", value))


def _looks_like_material_title(value: str) -> bool:
    value = _clean(value)
    if not (3 <= len(value) <= 120):
        return False
    if _HEADER_NOISE_RE.search(value):
        return False
    if re.search(r"\d{2,}|[@:]", value):
        return False
    if _CHEM_FORMULA_RE.match(value):
        return False
    letters = re.sub(r"[^A-Za-zА-Яа-я]", "", value)
    return len(letters) >= 3


def _material_from_top_lines(text: str) -> str | None:
    lines = [_clean(line) for line in text.splitlines()]
    candidates = [line for line in lines[:25] if _looks_like_material_title(line)]
    if not candidates:
        return None

    merged: list[str] = []
    for line in candidates[:4]:
        if merged and len(merged[-1]) + len(line) <= 90 and line.isupper():
            merged[-1] = f"{merged[-1]} {line}"
        else:
            merged.append(line)
    return merged[0] if merged else None


def _guess_sop_form(filename: str, text: str) -> str:
    hay = f"{filename}\n{text[:5000]}".lower()
    if re.search(r"(субстанц|сырь|spc[_\-\s]*суб|соп[\s–-]*533|sop[\s–-]*533)", hay):
        return "533"
    if re.search(r"(готов(ая|ой)?\s+продукц|гп\b|finished product|соп[\s–-]*548|sop[\s–-]*548|таблет|капсул)", hay):
        return "548"
    return "533"


def _infer_unit(spec: str) -> str | None:
    compact = _clean(spec)
    if len(compact) > 48:
        return None
    if not re.search(r"\d", compact):
        return None
    if "%" in compact:
        return "%"
    unit_match = _UNIT_RE.search(compact)
    return unit_match.group(1) if unit_match else None


def _extract_docx_text(data: bytes) -> tuple[str, list[list[str]]]:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - dependency is installed in Docker
        raise HTTPException(status_code=500, detail="python-docx is not available") from exc

    doc = Document(BytesIO(data))
    lines: list[str] = []
    rows: list[list[str]] = []
    for p in doc.paragraphs:
        text = _clean(p.text)
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
                lines.append(" | ".join(cells))
    return "\n".join(lines), rows


def _extract_pdf_text(data: bytes) -> tuple[str, list[list[str]]]:
    try:
        import pypdfium2 as pdfium
    except Exception as exc:  # pragma: no cover - dependency is installed in Docker
        raise HTTPException(status_code=500, detail="pypdfium2 is not available") from exc

    with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
        tmp.write(data)
        tmp.flush()
        pdf = pdfium.PdfDocument(tmp.name)
        lines: list[str] = []
        for page in pdf:
            textpage = page.get_textpage()
            text = _clean(textpage.get_text_range())
            if text:
                lines.extend(line.strip() for line in text.splitlines() if line.strip())
            textpage.close()
            page.close()
        pdf.close()
    return "\n".join(lines), []


def _extract_document(filename: str, data: bytes) -> tuple[str, list[list[str]]]:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _extract_docx_text(data)
    if lower.endswith(".pdf"):
        return _extract_pdf_text(data)
    raise HTTPException(status_code=400, detail="Only .docx and .pdf specification files are supported")


def _guess_header(filename: str, text: str) -> tuple[str, str | None, str]:
    nd_code = _code_from_filename(filename)
    if not nd_code:
        nd_match = _ND_RE.search(text)
        nd_code = _normalise_code(nd_match.group(0)) if nd_match else f"ND-IMPORT-{_normalise_code(_filename_stem(filename))[:40]}"
    rev_match = _REV_RE.search(text)
    revision = _clean(rev_match.group(1)) if rev_match else None

    material = ""
    for pattern in (
        r"(?:наименование\s+(?:сырья|материала|продукта)|material\s+name|product\s+name)\s*[:\-]\s*(.+)",
        r"(?:спецификация|specification)\s+(?:на|for)\s+(.+)",
    ):
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            material = _clean(m.group(1).splitlines()[0])
            break
    if material and _HEADER_NOISE_RE.search(material):
        material = ""
    filename_material = _material_from_filename(filename)
    if not material:
        top_material = _material_from_top_lines(text)
        material = filename_material if _has_cyrillic(filename_material) else top_material or filename_material or _filename_stem(filename)
    return nd_code[:128], revision[:32] if revision else None, material[:255]


def _row_to_param(cells: list[str]) -> dict | None:
    cleaned = [_clean(c) for c in cells if _clean(c)]
    if len(cleaned) < 2:
        return None
    joined = " ".join(cleaned).lower()
    if any(h in joined for h in ("показатель", "наименование показателя", "test", "parameter")) and any(
        h in joined for h in ("норма", "spec", "requirement", "треб")
    ):
        return None
    if cleaned[0].isdigit() and len(cleaned) >= 3:
        cleaned = cleaned[1:]
    name = cleaned[0]
    spec = cleaned[1]
    method = cleaned[2] if len(cleaned) >= 3 else None
    unit = cleaned[3] if len(cleaned) >= 4 else None
    if not unit:
        unit = _infer_unit(spec)
    if len(name) < 2 or len(spec) < 1:
        return None
    category = "microbiological" if re.search(r"(микро|бактер|дрож|плес|salmonella|e\.?\s*coli|cfu|кое)", name, re.IGNORECASE) else "physicochemical"
    return {
        "category": category,
        "parameter_name": name[:255],
        "specification": spec,
        "method_reference": method[:255] if method else None,
        "unit": unit[:32] if unit else None,
    }


def _parse_parameters(text: str, rows: list[list[str]]) -> list[dict]:
    params: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for row in rows:
        parsed = _row_to_param(row)
        if parsed:
            key = (parsed["parameter_name"].lower(), parsed["specification"].lower())
            if key not in seen:
                seen.add(key)
                params.append(parsed)

    if params:
        return params

    for line in text.splitlines():
        line = _clean(line)
        if not line or len(line) < 8:
            continue
        parts = [p.strip(" -:") for p in re.split(r"\s{2,}|\t|\|", line) if p.strip(" -:")]
        parsed = _row_to_param(parts)
        if parsed:
            key = (parsed["parameter_name"].lower(), parsed["specification"].lower())
            if key not in seen:
                seen.add(key)
                params.append(parsed)
    return params


def list_specifications(db: Session, include_inactive: bool = True) -> list[MaterialSpecification]:
    query = db.query(MaterialSpecification)
    if not include_inactive:
        query = query.filter(MaterialSpecification.is_active.is_(True))
    return query.order_by(MaterialSpecification.material_name).all()


def get_specification(db: Session, spec_id: UUID) -> MaterialSpecification:
    spec = db.get(MaterialSpecification, spec_id)
    if not spec:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Specification not found")
    return spec


def _replace_parameters(db: Session, spec: MaterialSpecification, payload: MaterialSpecificationInput) -> None:
    db.query(SpecificationParameter).filter(SpecificationParameter.specification_id == spec.id).delete()
    db.flush()
    for idx, p in enumerate(payload.parameters, start=1):
        db.add(
            SpecificationParameter(
                specification_id=spec.id,
                category=p.category,
                ordinal=idx,
                parameter_name=p.parameter_name,
                specification=p.specification,
                method_reference=p.method_reference,
                unit=p.unit,
            )
        )


def create_specification(db: Session, user: CurrentUser, payload: MaterialSpecificationInput) -> MaterialSpecification:
    require_permission(user, "MANAGE_SPECIFICATIONS")
    spec = MaterialSpecification(
        nd_code=payload.nd_code,
        revision=payload.revision,
        material_name=payload.material_name,
        material_id=payload.material_id,
        match_keywords=payload.match_keywords,
        sop_form=payload.sop_form,
        micro_required=payload.micro_required,
        micro_method_ref=payload.micro_method_ref,
        is_active=payload.is_active,
        effective_date=payload.effective_date,
        notes=payload.notes,
        created_by=user.id,
    )
    db.add(spec)
    db.flush()
    _replace_parameters(db, spec, payload)
    write_audit(
        db, user, object_type="material_specification", object_id=str(spec.id),
        action_type="CREATE_SPECIFICATION",
        new_value={"nd_code": spec.nd_code, "material": spec.material_name, "parameters": len(payload.parameters)},
    )
    db.commit()
    db.refresh(spec)
    return spec


def import_specification_document(db: Session, user: CurrentUser, filename: str, data: bytes) -> MaterialSpecification:
    """Create a review draft from an uploaded ND/specification document."""
    require_permission(user, "MANAGE_SPECIFICATIONS")
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    text, rows = _extract_document(filename, data)
    if len(_clean(text)) < 20:
        raise HTTPException(status_code=422, detail="No readable text found in the document")

    nd_code, revision, material_name = _guess_header(filename, text)
    existing_codes = {r[0] for r in db.query(MaterialSpecification.nd_code).all()}
    base_code = nd_code
    suffix = 2
    while nd_code in existing_codes:
        nd_code = f"{base_code[:112]}-IMP-{suffix}"
        suffix += 1

    params = _parse_parameters(text, rows)
    micro_required = any(p["category"] == "microbiological" for p in params)
    payload = MaterialSpecificationInput(
        nd_code=nd_code,
        revision=revision,
        material_name=material_name,
        material_id=None,
        match_keywords=material_name,
        sop_form=_guess_sop_form(filename, text),
        micro_required=micro_required,
        micro_method_ref=None,
        is_active=False,
        effective_date=None,
        notes=(
            f"Черновик импортирован из файла {filename}. "
            f"Проверьте распознанные поля и параметры перед вводом в действие."
        ),
        parameters=params,
    )
    spec = create_specification(db, user, payload)
    write_audit(
        db, user, object_type="material_specification", object_id=str(spec.id),
        action_type="IMPORT_SPECIFICATION_DOCUMENT",
        new_value={"filename": filename, "parameters": len(params), "draft": True},
    )
    db.commit()
    db.refresh(spec)
    return spec


def update_specification(db: Session, user: CurrentUser, spec_id: UUID, payload: MaterialSpecificationInput) -> MaterialSpecification:
    require_permission(user, "MANAGE_SPECIFICATIONS")
    spec = get_specification(db, spec_id)
    spec.nd_code = payload.nd_code
    spec.revision = payload.revision
    spec.material_name = payload.material_name
    spec.material_id = payload.material_id
    spec.match_keywords = payload.match_keywords
    spec.sop_form = payload.sop_form
    spec.micro_required = payload.micro_required
    spec.micro_method_ref = payload.micro_method_ref
    spec.is_active = payload.is_active
    spec.effective_date = payload.effective_date
    spec.notes = payload.notes
    _replace_parameters(db, spec, payload)
    write_audit(
        db, user, object_type="material_specification", object_id=str(spec.id),
        action_type="UPDATE_SPECIFICATION",
        new_value={"nd_code": spec.nd_code, "material": spec.material_name, "parameters": len(payload.parameters)},
    )
    db.commit()
    db.refresh(spec)
    return spec


def delete_specification(db: Session, user: CurrentUser, spec_id: UUID) -> None:
    require_permission(user, "MANAGE_SPECIFICATIONS")
    spec = get_specification(db, spec_id)
    write_audit(
        db, user, object_type="material_specification", object_id=str(spec.id),
        action_type="DELETE_SPECIFICATION",
        old_value={"nd_code": spec.nd_code, "material": spec.material_name},
    )
    db.delete(spec)
    db.commit()


def resolve_for_lot(db: Session, lot: Lot) -> MaterialSpecification | None:
    """Подбирает активную спецификацию для партии.

    Приоритет: точное совпадение по material_id, затем совпадение по
    ключевым словам/наименованию материала (без учёта регистра).
    """
    active = (
        db.query(MaterialSpecification)
        .filter(MaterialSpecification.is_active.is_(True))
        .all()
    )
    if not active:
        return None

    # 1. По material_id.
    if lot.material_id:
        for spec in active:
            if spec.material_id and spec.material_id == lot.material_id:
                return spec

    # 2. По ключевым словам / наименованию материала.
    material = db.get(Material, lot.material_id) if lot.material_id else None
    hay = " ".join(filter(None, [
        (material.name if material else "") or "",
        (material.code if material else "") or "",
    ])).lower()
    if not hay.strip():
        return None
    for spec in active:
        tokens: list[str] = []
        if spec.match_keywords:
            tokens.extend(t.strip().lower() for t in spec.match_keywords.replace(",", " ").split())
        if spec.material_name:
            tokens.append(spec.material_name.lower())
        for token in tokens:
            if token and (token in hay or hay in token):
                return spec
    return None
