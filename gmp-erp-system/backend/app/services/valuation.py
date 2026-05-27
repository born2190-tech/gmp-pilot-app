"""Стоимостная оценка запасов по счетам учёта (оборотная ведомость).

Метод — партионная себестоимость: стоимость партии = остаток × unit_cost.
Баланс счёта = Σ стоимостей партий на этом счёте. Обороты — из движений,
оценённых по себестоимости партии (приход на счёт / расход со счёта).
"""
from __future__ import annotations

from datetime import date

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.inventory import InventoryMovement, Lot
from app.models.master_data import InventoryAccount, Warehouse


def account_ledger(db: Session, date_from: date | None = None, date_to: date | None = None) -> dict:
    accounts = db.query(InventoryAccount).order_by(InventoryAccount.code).all()
    acc_by_id = {a.id: a for a in accounts}

    # ── Текущие балансы: Σ(остаток × себестоимость) по счёту ──
    value_expr = func.coalesce(Lot.quantity * Lot.unit_cost, 0.0)
    bal_rows = (
        db.query(
            Lot.account_id.label("account_id"),
            func.count(Lot.id).label("lots_count"),
            func.coalesce(func.sum(value_expr), 0.0).label("balance_value"),
        )
        .filter(Lot.quantity > 0)
        .group_by(Lot.account_id)
        .all()
    )
    balances = {r.account_id: r for r in bal_rows}

    # Разбивка балансов по складам (тип склада) внутри счёта.
    by_wh_rows = (
        db.query(
            Lot.account_id.label("account_id"),
            Warehouse.warehouse_type.label("warehouse_type"),
            func.coalesce(func.sum(value_expr), 0.0).label("value"),
        )
        .join(Warehouse, Warehouse.id == Lot.warehouse_id)
        .filter(Lot.quantity > 0)
        .group_by(Lot.account_id, Warehouse.warehouse_type)
        .all()
    )
    by_wh: dict = {}
    for r in by_wh_rows:
        by_wh.setdefault(r.account_id, []).append({"warehouse_type": r.warehouse_type, "value": float(r.value or 0)})

    # ── Обороты за период: движения, оценённые по себестоимости партии ──
    move_value = func.abs(InventoryMovement.quantity_delta) * func.coalesce(Lot.unit_cost, 0.0)
    mq = db.query(
        InventoryMovement.from_account_id.label("from_account_id"),
        InventoryMovement.to_account_id.label("to_account_id"),
        move_value.label("value"),
    ).join(Lot, Lot.id == InventoryMovement.lot_id)
    if date_from:
        mq = mq.filter(func.date(InventoryMovement.created_at) >= date_from)
    if date_to:
        mq = mq.filter(func.date(InventoryMovement.created_at) <= date_to)
    in_value: dict = {}
    out_value: dict = {}
    for r in mq.all():
        v = float(r.value or 0)
        if r.to_account_id is not None:
            in_value[r.to_account_id] = in_value.get(r.to_account_id, 0.0) + v
        if r.from_account_id is not None:
            out_value[r.from_account_id] = out_value.get(r.from_account_id, 0.0) + v

    items = []
    total_balance = 0.0
    for acc in accounts:
        bal = balances.get(acc.id)
        balance_value = float(bal.balance_value) if bal else 0.0
        lots_count = int(bal.lots_count) if bal else 0
        total_balance += balance_value
        items.append({
            "account_id": str(acc.id),
            "account_code": acc.code,
            "account_name": acc.name,
            "account_group": acc.account_group,
            "lots_count": lots_count,
            "balance_value": balance_value,
            "in_value": round(in_value.get(acc.id, 0.0), 2),
            "out_value": round(out_value.get(acc.id, 0.0), 2),
            "by_warehouse": by_wh.get(acc.id, []),
        })

    # Партии без счёта (исторические/непривязанные) — отдельной строкой.
    unassigned = balances.get(None)
    if unassigned and float(unassigned.balance_value) > 0:
        items.append({
            "account_id": None,
            "account_code": "—",
            "account_name": "Без счёта учёта",
            "account_group": None,
            "lots_count": int(unassigned.lots_count),
            "balance_value": float(unassigned.balance_value),
            "in_value": 0.0,
            "out_value": 0.0,
            "by_warehouse": by_wh.get(None, []),
        })
        total_balance += float(unassigned.balance_value)

    return {"currency": "UZS", "total_balance": round(total_balance, 2), "accounts": items}


def export_account_ledger_xlsx(db: Session, date_from: date | None = None, date_to: date | None = None) -> bytes:
    """Оборотная ведомость по счетам → XLSX."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = account_ledger(db, date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "Оборотная ведомость"

    period = ""
    if date_from or date_to:
        period = f" за период {date_from or '…'} — {date_to or '…'}"
    ws["A1"] = f"Оборотная ведомость по счетам учёта{period}"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:F1")

    headers = ["Счёт", "Наименование", "Партий", "Приход (∑)", "Расход (∑)", "Баланс"]
    header_row = 3
    for col, title in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="334155")
        c.alignment = Alignment(horizontal="center")

    r = header_row + 1
    for a in data["accounts"]:
        ws.cell(row=r, column=1, value=a["account_code"])
        ws.cell(row=r, column=2, value=a["account_name"])
        ws.cell(row=r, column=3, value=a["lots_count"])
        ws.cell(row=r, column=4, value=a["in_value"])
        ws.cell(row=r, column=5, value=a["out_value"])
        ws.cell(row=r, column=6, value=a["balance_value"])
        r += 1

    total_cell = ws.cell(row=r, column=2, value="ИТОГО")
    total_cell.font = Font(bold=True)
    bal = ws.cell(row=r, column=6, value=round(data["total_balance"], 2))
    bal.font = Font(bold=True)

    widths = [14, 40, 10, 16, 16, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row + 1, max_row=r, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
